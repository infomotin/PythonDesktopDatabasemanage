import csv, io, re, json, hashlib, tempfile, threading, shutil, time, datetime, warnings, random, string, uuid
import numpy as np
import pandas as pd
from pathlib import Path
from django.conf import settings
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Count, Sum, F, Q
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_http_methods
from django.core.cache import cache
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.core.files.storage import default_storage


def read_excel_file(file_path):
    return pd.read_excel(file_path)


def generate_excel_file(df, sheet_name="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    return wb


def clean_col_names(df):
    df.columns = [str(c).strip().replace(" ", "_").lower() for c in df.columns]
    return df


def df_to_insert_sql(df, table_name):
    import re
    cols = list(df.columns)
    safe = [re.sub(r"[^a-zA-Z0-9_]", "_", c) for c in cols]
    statements = []
    for _, row in df.iterrows():
        vals = []
        for v in row:
            if isinstance(v, str):
                vals.append("'" + v.replace("'", "''") + "'")
            elif v is None:
                vals.append("NULL")
            else:
                vals.append(str(v))
        statements.append(f"INSERT INTO {table_name} ({', '.join(safe)}) VALUES ({', '.join(vals)});")
    return "\n".join(statements)


def validate_dataframe(df, required=None):
    errors = []
    for col in (required or []):
        if col not in df.columns:
            errors.append(f"Missing required column: {col}")
    return errors
