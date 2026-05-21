import json, csv, io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def read_excel_to_dataframe(file_path, sheets=None, **kwargs):
    path = str(file_path)
    try:
        if isinstance(sheets, list):
            return [pd.read_excel(path, sheet_name=s, **kwargs) for s in sheets]
        return pd.read_excel(path, sheet_name=sheets or 0, **kwargs)
    except Exception as exc:
        raise Exception(f"Failed to read Excel: {exc}")


def clean_column_names(df):
    df.columns = [str(c).strip().replace(" ", "_").lower() for c in df.columns]
    return df


def coerce_types(df, mapping=None):
    for col, dtype in (mapping or {}).items():
        if col in df.columns:
            try:
                df[col] = df[col].astype(dtype)
            except Exception:
                pass
    return df


def validate_excel_import(df, required_cols=None, rules=None):
    errors = []
    required_cols = required_cols or []
    for col in required_cols:
        if col not in df.columns:
            errors.append(f"Missing required column: {col}")
    return errors


def generate_excel(df, sheet_name="Sheet1", filename="export.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
    for col_idx, col_name in enumerate(df.columns, 1):
        max_len = max(len(str(col_name)), df[col_name].astype(str).str.len().max() if len(df) else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)
    return wb


def generate_csv(df, filename="export.csv"):
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    return buf.getvalue()


def generate_excel_multisheet(dfs, sheet_names=None, filename="multi.xlsx"):
    wb = Workbook()
    wb.remove(wb.active)
    for i, df in enumerate(dfs):
        name = sheet_names[i] if sheet_names and i < len(sheet_names) else f"Sheet{i+1}"
        ws = wb.create_sheet(title=name)
        for col_idx, col_name in enumerate(df.columns, 1):
            ws.cell(row=1, column=col_idx, value=col_name)
        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            for col_idx, val in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)
    return wb


def df_to_insert_sql(df, table_name):
    import re
    columns = list(df.columns)
    safe_cols = [re.sub(r"[^a-zA-Z0-9_]", "_", c) for c in columns]
    statements = []
    for _, row in df.iterrows():
        vals = []
        for v in row:
            if isinstance(v, str):
                vals.append("'" + v.replace("'", "''") + "'")
            elif v is None:
                vals.append("NULL")
            else:
                vals.append(str(v))
        stmt = f"INSERT INTO {table_name} ({', '.join(safe_cols)}) VALUES ({', '.join(vals)});"
        statements.append(stmt)
    return "\n".join(statements)


class DataImportExportError(Exception):
    pass
